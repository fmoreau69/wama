"""La brique de traduction (`common/utils/translator.py`) — moteur de l'app Translator (2026-10-16).

Elle n'avait AUCUN test alors que le pipeline de prompts de toutes les apps et l'assistant
l'appellent. Ce qui est tenu : le passthrough, le modèle désigné par une CLÉ DE CATALOGUE (voie
locale gouvernée pour `ollama:*`, `chat_with_catalog_model` sinon, avec la clé de l'utilisateur),
l'ancien nom nu accepté, l'aperçu au fil des morceaux, et un cache qui distingue les modèles.
"""
import shutil
from pathlib import Path
from unittest import mock

from django.test import SimpleTestCase, TestCase

from wama.common.utils import llm_utils, translator


class TranslatorServiceTest(SimpleTestCase):

    def _service(self, **kw):
        return translator.TranslatorService(use_cache=False, **kw)

    def test_same_language_or_empty_text_calls_no_model(self):
        with mock.patch.object(llm_utils, 'llm_chat') as local, \
                mock.patch.object(llm_utils, 'chat_with_catalog_model') as remote:
            self.assertEqual('Bonjour', self._service().translate('Bonjour', 'fr', 'fr')['text'])
            self.assertTrue(self._service().translate('  ', 'fr', 'en')['ok'])
        local.assert_not_called()
        remote.assert_not_called()

    def test_a_local_key_takes_the_governed_ollama_path(self):
        with mock.patch.object(llm_utils, 'llm_chat', return_value=('Hello', None)) as local:
            out = self._service().translate('Bonjour', 'fr', 'en')
        self.assertEqual('Hello', out['text'])
        self.assertEqual(('ollama', 'translategemma:12b'),
                         (local.call_args.kwargs['provider'], local.call_args.kwargs['model']))

    def test_a_remote_key_goes_through_the_catalog_call_with_the_user(self):
        user = object()
        with mock.patch.object(llm_utils, 'chat_with_catalog_model',
                               return_value=('Hello', None)) as remote, \
                mock.patch.object(llm_utils, 'llm_chat') as local:
            out = self._service(model='ilaas:qwen-3.8-27b', user=user).translate(
                'Bonjour', 'fr', 'en')
        self.assertEqual('Hello', out['text'])
        self.assertEqual('ilaas:qwen-3.8-27b', remote.call_args.args[0])
        self.assertIs(user, remote.call_args.kwargs['user'])
        local.assert_not_called()

    def test_the_former_bare_ollama_name_is_still_accepted(self):
        self.assertEqual('ollama:translategemma:12b',
                         self._service(model='translategemma:12b').model)

    def test_each_translated_chunk_is_reported_as_it_comes(self):
        text = 'a' * 3000 + '\n\n' + 'b' * 3000          # deux morceaux (> _CHUNK_CHARS)
        seen = []
        with mock.patch.object(llm_utils, 'llm_chat', side_effect=[('A', None), ('B', None)]):
            out = self._service().translate(text, 'fr', 'en', on_chunk=seen.append)
        self.assertEqual(['A', 'A\n\nB'], seen)
        self.assertEqual('A\n\nB', out['text'])

    def test_an_unknown_source_is_left_to_the_model_and_never_named_auto(self):
        """Mesuré le 2026-10-16 sur translategemma : « de « auto » vers « en » » rendait le texte
        INCHANGÉ. La consigne nomme les langues par code ISO et fait détecter la source."""
        with mock.patch.object(llm_utils, 'llm_chat', return_value=('Hello', None)) as local:
            self._service().translate('Bonjour', 'auto', 'en')
            unknown = local.call_args.kwargs['messages'][0]['content']
            self._service().translate('Bonjour', 'fr', 'en')
            known = local.call_args.kwargs['messages'][0]['content']
        self.assertNotIn('« auto »', unknown)
        self.assertIn('détecte', unknown)
        self.assertIn('ISO 639-1 « en »', unknown)
        self.assertIn('« fr »', known)                    # contre-épreuve : la source connue se dit

    def test_only_a_marked_text_is_told_to_keep_its_tags(self):
        with mock.patch.object(llm_utils, 'llm_chat', return_value=('Hello', None)) as local:
            self._service().translate('Voir la <t1>note</t1>.', 'fr', 'en')
            marked = local.call_args.kwargs['messages'][0]['content']
            self._service().translate('Voir la note.', 'fr', 'en')
            bare = local.call_args.kwargs['messages'][0]['content']
        self.assertIn('balises <t1>', marked)
        self.assertNotIn('balises', bare)              # a bare text keeps its prompt, and its cache

    def test_a_rewritten_prompt_does_not_serve_the_old_cached_translations(self):
        key = translator.TranslatorService._cache_key
        before = key('Bonjour', 'auto', 'en', [], 'ollama:translategemma:12b')
        with mock.patch.object(translator, 'PROMPT_VERSION', translator.PROMPT_VERSION + 1):
            self.assertNotEqual(before, key('Bonjour', 'auto', 'en', [], 'ollama:translategemma:12b'))

    def test_two_models_never_share_a_cached_translation(self):
        key = translator.TranslatorService._cache_key
        self.assertNotEqual(key('Bonjour', 'fr', 'en', [], 'ollama:translategemma:12b'),
                            key('Bonjour', 'fr', 'en', [], 'ilaas:qwen-3.8-27b'))


def _fake_translate(self, text, source_lang, target_lang, glossary=None, timeout=120,
                    on_chunk=None):
    """A deterministic « translation » : the text upper-cased, with a Polish letter (that cp1252
    cannot carry) — enough to see WHERE each block landed, and that the form stayed."""
    return {'ok': True, 'text': f'Ł {text.upper()}', 'cached': False}


class TranslateDocumentTest(SimpleTestCase):
    """2026-10-28 — a document translated KEEPS ITS FORM : same format, each block translated in
    its place (`document_text`, unit `block`), never the plain text rewritten as a `.md`."""

    def setUp(self):
        import shutil
        import tempfile
        from pathlib import Path
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, True)
        patcher = mock.patch.object(translator.TranslatorService, 'translate', _fake_translate)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_a_docx_keeps_its_paragraph_form_and_its_format(self):
        import docx
        source = self.tmp / 'in.docx'
        document = docx.Document()
        title = document.add_paragraph()
        title.add_run('Rapport annuel').bold = True
        document.add_paragraph('')
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Dans le tableau'
        document.save(str(source))
        seen = []
        out = translator.translate_document(source, self.tmp / 'out.docx', 'fr', 'en',
                                            on_chunk=seen.append)
        result = docx.Document(out)
        self.assertEqual('Ł RAPPORT ANNUEL', result.paragraphs[0].text)
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        self.assertEqual('', result.paragraphs[1].text)            # an empty paragraph stays
        self.assertEqual('Ł DANS LE TABLEAU', result.tables[0].cell(0, 0).text)
        self.assertEqual('Ł RAPPORT ANNUEL\n\nŁ DANS LE TABLEAU', seen[-1])

    def test_a_pdf_is_translated_block_by_block_and_keeps_its_drawings(self):
        import fitz
        source = self.tmp / 'in.pdf'
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.draw_rect(fitz.Rect(40, 40, 300, 200), color=(1, 0, 0))
            page.insert_text((50, 80), 'Premier bloc', fontsize=12)
            page.insert_text((50, 400), 'Second bloc', fontsize=12)
            pdf.save(str(source))
        out = translator.translate_document(source, self.tmp / 'out.pdf', 'fr', 'en')
        with fitz.open(out) as pdf:
            text = pdf[0].get_text()
            drawings = pdf[0].get_drawings()
        self.assertIn('PREMIER BLOC', text)
        self.assertIn('SECOND BLOC', text)
        self.assertNotIn('Premier bloc', text)                 # the original is gone, not hidden
        self.assertTrue(drawings, 'the drawn rectangle must survive the rewrite')

    def test_a_text_file_its_encoding_cannot_carry_passes_to_utf8(self):
        source = self.tmp / 'in.txt'
        source.write_bytes('Été'.encode('cp1252'))
        out = translator.translate_document(source, self.tmp / 'out.txt', 'fr', 'pl')
        from pathlib import Path
        self.assertEqual('Ł ÉTÉ', Path(out).read_text(encoding='utf-8'))

    def test_a_document_without_text_is_refused_with_the_reason(self):
        import fitz
        source = self.tmp / 'scan.pdf'
        with fitz.open() as pdf:
            pdf.new_page()
            pdf.save(str(source))
        with self.assertRaisesRegex(ValueError, 'PDF scanné'):
            translator.translate_document(source, self.tmp / 'out.pdf', 'fr', 'en')

    def test_a_bold_word_inside_a_paragraph_stays_bold_and_the_preview_has_no_tags(self):
        """2026-10-29 — the block leaves MARKED (`<t1>…</t1>`), the rewrite gives the bold back to
        the translated word ; before, the whole paragraph took the format of its first run."""
        import docx
        source = self.tmp / 'in.docx'
        document = docx.Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('Voir la ')
        paragraph.add_run('note').bold = True
        paragraph.add_run(' plus bas.')
        document.save(str(source))
        seen = []
        out = translator.translate_document(source, self.tmp / 'out.docx', 'fr', 'en',
                                            on_chunk=seen.append)
        runs = [(r.text, bool(r.bold)) for r in docx.Document(out).paragraphs[0].runs if r.text]
        self.assertEqual([('Ł VOIR LA ', False), ('NOTE', True), (' PLUS BAS.', False)], runs)
        self.assertEqual('Ł VOIR LA NOTE PLUS BAS.', seen[-1])

    def test_a_label_gets_no_final_stop_and_an_axis_letter_no_paragraph(self):
        """2026-10-29, figure of a real paper : the model added « . » to every label, and answered
        a paragraph to the axis letter « R »."""
        import docx
        source = self.tmp / 'in.docx'
        document = docx.Document()
        for text in ('Moyenne spatiale', 'R', 'Une phrase.', 'Signal'):
            document.add_paragraph(text)
        document.save(str(source))
        answers = {'Moyenne spatiale': 'Spatial averaging.', 'Une phrase.': 'A sentence.',
                   'Signal': 'The word signal designates, in this figure, the measured quantity.'}
        sent = []

        def model(self, text, *a, **k):
            sent.append(text)
            return {'ok': True, 'text': answers[text], 'cached': False}

        with mock.patch.object(translator.TranslatorService, 'translate', model):
            out = translator.translate_document(source, self.tmp / 'out.docx', 'fr', 'en')
        self.assertNotIn('R', sent)                                    # one letter : not sent
        self.assertEqual(['Spatial averaging', 'R', 'A sentence.', 'Signal'],
                         [p.text for p in docx.Document(out).paragraphs])

    def test_a_glossary_written_one_term_per_line_gives_terms_not_characters(self):
        with mock.patch.object(translator.TranslatorService, 'translate',
                               return_value={'ok': True, 'text': 'x', 'cached': False}) as call:
            translator.translate_text('Bonjour', 'fr', 'en', glossary='WAMA\n\n Lescot \n')
        self.assertEqual(['WAMA', 'Lescot'], call.call_args.kwargs['glossary'])

    def test_a_failed_translation_raises_its_reason(self):
        failing = lambda *a, **k: {'ok': False, 'error': 'modèle absent', 'cached': False}  # noqa: E731
        with mock.patch.object(translator.TranslatorService, 'translate', failing), \
                self.assertRaisesRegex(RuntimeError, 'modèle absent'):
            translator.translate_text('Bonjour', 'fr', 'en')


class TranslationIsDeclaredNotWrittenTest(SimpleTestCase):
    """The translation SPECIALITY is declared on the schema (`options_resolution`), read by the
    draw and by the `codegen` role — never passed by hand in a glue."""

    def test_the_declared_speciality_reaches_the_draw_as_a_single_value(self):
        from wama.common.utils import auto_model
        self.assertIn('specialization', auto_model.RESOLUTION_KEYS)
        self.assertEqual('translation', auto_model.read_refinement('specialization', 'translation'))
        with mock.patch.object(auto_model, 'catalog_field', return_value={
                'options_resolution': {'specialization': 'translation', 'unknown': 'x'}}):
            self.assertEqual({'specialization': 'translation'},
                             auto_model.declared_resolution('translator_01'))

    def test_codegen_shows_the_translation_brick_instead_of_the_bare_llm_call(self):
        from wama.common.tests.tests_dev_ai_bridge import _charger
        codegen = _charger('run_codegen')
        llm_model = [{'manifest_kind': 'model', 'body': {'identity': {'model_type': 'llm'}}}]
        select = {'name': 'model', 'options_resolution': {'specialization': 'translation'}}
        man = {'body': {'params': {'schemas': {'P': [select]}}}}
        bricks = [names for _key, names in codegen._triggered(llm_model, man)]
        self.assertIn(('translate_text', 'translate_document'), bricks)
        self.assertNotIn(('chat_with_catalog_model',), bricks)
        # Counter-proof : without the declared speciality, the bare LLM call is what it sees.
        plain = {'body': {'params': {'schemas': {'P': [{'name': 'model'}]}}}}
        self.assertEqual([('chat_with_catalog_model',)],
                         [names for _key, names in codegen._triggered(llm_model, plain)])
        self.assertEqual({'text', 'source_language', 'target_language', 'model', 'user',
                          'glossary', 'on_chunk', 'timeout'},
                         codegen.signatures_briques()['translate_text'])

    def test_codegen_shows_the_ocr_brick_when_a_setting_follows_the_input_s_text_layer(self):
        """2026-10-29 : a switch proposed by `needs_ocr` (`Param.default_from`) means an « OCR of
        scanned pages » process — its glue is the common brick, which the role must SEE."""
        from wama.common.tests.tests_dev_ai_bridge import _charger
        codegen = _charger('run_codegen')
        switch = {'name': 'ocr_scanned_pages', 'type': 'toggle',
                  'default_from': {'port': 'work_file', 'property': 'needs_ocr'}}
        man = {'body': {'params': {'schemas': {'P': [switch]}}}}
        self.assertIn(('read_ocr_file',), [names for _key, names in codegen._triggered([], man)])
        plain = {'body': {'params': {'schemas': {'P': [dict(switch, default_from=None)]}}}}
        self.assertEqual([], codegen._triggered([], plain))
        self.assertEqual({'item', 'field'}, codegen.signatures_briques()['read_ocr_file'])

    def test_an_app_with_the_live_preview_is_shown_how_to_publish_it(self):
        from wama.common.tests.tests_dev_ai_bridge import _charger
        codegen = _charger('run_codegen')
        live = {'body': {'capabilities': {'during_preview': True}}}
        self.assertIn(('publish_partial_text', 'clear_partial'),
                      [names for _key, names in codegen._triggered([], live)])
        self.assertEqual([], codegen._triggered([], {'body': {'capabilities': {'during_preview': False}}}))
        self.assertEqual({'app_name', 'pk', 'text'}, codegen.signatures_briques()['publish_partial_text'])

    def test_an_upstream_output_read_by_a_brick_by_its_name_is_read(self):
        """2026-10-30 : the translate glue relied on `read_ocr_file(item, 'ocr_file')`; the judge
        called the OCR output « never read » — the instrument was wrong, as for `getattr` before."""
        from wama.common.tests.tests_dev_ai_bridge import _charger
        codegen = _charger('run_codegen')
        ocr = {'key': 'ocr', 'outputs': ['ocr_file'], 'depends_on': [], 'watched': []}
        translate = {'key': 'translate', 'outputs': ['output_file'], 'depends_on': ['ocr'],
                     'watched': []}
        glue = ("def _g(item, ctx):\n"
                "    lines = read_ocr_file(item, 'ocr_file')\n"
                "    return {'fields': {'output_file': 'x'}, 'output_ref': 'x'}\n")
        self.assertEqual([], codegen.process_checks(glue, '_g', translate, [ocr, translate]))
        ignored = glue.replace("read_ocr_file(item, 'ocr_file')", 'None')
        self.assertEqual(['sortie d\'amont jamais lue : item.ocr_file'],
                         codegen.process_checks(ignored, '_g', translate, [ocr, translate]))

    def test_a_common_process_whose_arguments_are_declared_is_rendered_not_left_as_a_hole(self):
        """The OCR process is COMMON (`pipeline_decl.COMMON_PROCESSES`, with the « Sortie »): its
        glue is RENDERED by the generator — the brick only needs the item's input field and the
        process's declared output. Three proposals of the role copied or invented the brick."""
        import json
        from pathlib import Path
        from django.conf import settings
        from wama.common.manifests.codegen.tasks_gen import render_tasks
        base = Path(settings.BASE_DIR) / 'manifests' / 'app_drafts'
        manifest = json.loads((base / 'translator.json').read_text(encoding='utf-8'))
        source, _why = render_tasks(manifest)
        ocr_glue = source.split('def _process_translate_text_task_ocr')[1].split('\ndef ')[0]
        self.assertIn("return ocr_step('source_document', 'ocr_file')(item, ctx)", ocr_glue)
        self.assertNotIn('NotImplementedError', ocr_glue)
        translate_hole = source.split('def _process_translate_text_task_translate')[1]
        self.assertIn('NotImplementedError', translate_hole.split('\ndef ')[0],
                      'translation is not common: its glue stays with the role')
        compile(source, 'tasks.py', 'exec')

    def test_the_judge_sees_what_the_first_generated_glue_invented(self):
        """2026-10-28 : the first glue generated for the Translator passed the checks with ONE
        warning out of five defects — `fields` built in a variable, an invented item attribute,
        `ctx.output_place` unpacked into three values."""
        from wama.common.tests.tests_dev_ai_bridge import _charger
        codegen = _charger('run_codegen')
        glue = ('def _process_t(item, ctx):\n'
                '    from wama.common.utils.document_text import ensure_text_file\n'
                '    text = item.options\n'
                '    rel, folder, name = ctx.output_place("txt", source_name=item.filename)\n'
                '    ctx.progress(5, colour="red")\n'
                '    fields = {"result_text": text, "model": "x"}\n'
                '    return {"fields": fields}\n')
        warnings = '\n'.join(codegen.controles(
            glue, '_process_t', app_id='translator_01',
            item_fields=['prompt', 'model', 'output_file'],
            settings_fields=['prompt', 'model'], item_properties=['filename'])['warnings'])
        for expected in ("['result_text']", "['options']", "réécrit un RÉGLAGE de l'utilisateur",
                         'DEUX valeurs', 'ctx.progress', "['ensure_text_file']"):
            self.assertIn(expected, warnings)
        # Counter-proof : a glue that keeps to the item and the context passes.
        sound = ('def _process_t(item, ctx):\n'
                 '    from wama.common.utils.translator import translate_document\n'
                 '    rel, path = ctx.output_place("txt", source_name=item.filename)\n'
                 '    return {"fields": {"output_file": rel}, "models": [item.model]}\n')
        self.assertEqual([], codegen.controles(
            sound, '_process_t', app_id='translator_01',
            item_fields=['prompt', 'model', 'output_file'],
            settings_fields=['prompt', 'model'], item_properties=['filename'])['warnings'])

    def test_the_judge_sees_a_mode_that_returns_nothing_and_an_emptied_input(self):
        """2026-10-28 : two task glues passed the judge while returning `'output_file': None` in
        text mode and `'source_document': None` in document mode — the declared-output check only
        ran for a pipeline of several processes."""
        from wama.common.tests.tests_dev_ai_bridge import _charger
        codegen = _charger('run_codegen')
        glue = ('def _process_t(item, ctx):\n'
                '    rel, path = ctx.output_place("txt", source_name="x")\n'
                '    if item.prompt:\n'
                '        return {"fields": {"output_file": None}}\n'
                '    return {"fields": {"source_document": None, "output_file": rel}}\n')
        warnings = '\n'.join(codegen.result_checks(glue, '_process_t', 'output_file',
                                                   'source_document'))
        self.assertIn("1 retour(s) sur 2 n'écrivent pas la sortie déclarée `output_file`", warnings)
        self.assertIn("réécrit l'ENTRÉE de l'élément : source_document", warnings)
        # Counter-proof : every path writes the declared output, the input is only read.
        sound = ('def _process_t(item, ctx):\n'
                 '    rel, path = ctx.output_place("txt", source_name="x")\n'
                 '    fields = {"output_file": rel}\n'
                 '    if item.prompt:\n'
                 '        return {"fields": fields}\n'
                 '    return {"fields": {"output_file": rel}}\n')
        self.assertEqual([], codegen.result_checks(sound, '_process_t', 'output_file',
                                                   'source_document'))
        # 2026-10-29 : the absolute path thrown away — the returned file is never written.
        thrown = ('def _process_t(item, ctx):\n'
                  '    rel, _ = ctx.output_place("txt", source_name="x")\n'
                  '    return {"fields": {"output_file": rel}}\n')
        self.assertIn("chemin absolu d'`ctx.output_place` est jeté",
                      '\n'.join(codegen.result_checks(thrown, '_process_t', 'output_file')))

    def test_the_matter_shows_what_an_item_property_returns_not_only_its_name(self):
        """2026-10-29 : a glue tested `item.filename is None` to pick the text mode — the property
        returns '' without a document. Its name alone did not say so."""
        from wama.common.tests.tests_dev_ai_bridge import _charger
        codegen = _charger('run_codegen')
        _fields, props = codegen.champs_item('ReadingItem', 'reader')
        self.assertIn('filename', props)
        bodies = codegen.property_bodies('ReadingItem', 'reader', ['filename'])
        self.assertIn('def filename(self):', bodies)
        self.assertIn('original_filename', bodies)
        # Counter-proof : an unknown model gives no matter, never an error.
        self.assertEqual('', codegen.property_bodies('NoSuchItem', 'reader', ['filename']))

    def test_the_judge_sees_an_invented_draw_domain(self):
        """2026-10-29 : two glues drew in `{'source': 'catalog', 'specialization': …}` while the
        selector declares `{'model_type': 'llm'}`."""
        from wama.common.tests.tests_dev_ai_bridge import _charger
        codegen = _charger('run_codegen')
        man = {'body': {'params': {'schemas': {'P': [
            {'name': 'model', 'options_source': 'catalog', 'options_query': {'model_type': 'llm'},
             'options_resolution': {'specialization': 'translation'}}]}}}}
        invented = ('def _process_t(item, ctx):\n'
                    '    key = resolve_model_choice(item.model, spec={"source": "catalog", '
                    '"specialization": "translation"}, item=item)\n')
        self.assertIn("n'est pas un domaine déclaré",
                      '\n'.join(codegen.draw_domain_checks(invented, man)))
        # Counter-proof : the declared domain, or one held in a variable, passes.
        for spec in ('{"model_type": "llm"}', 'domain'):
            sound = ('def _process_t(item, ctx):\n'
                     f'    key = resolve_model_choice(item.model, spec={spec}, item=item)\n')
            self.assertEqual([], codegen.draw_domain_checks(sound, man))

    def test_a_task_body_reads_its_settings_and_draws_its_model_like_every_app(self):
        """2026-10-29 (Fabien) : settings and draw domain « parfaitement alignés sur le
        fonctionnement des autres apps » — one preamble, composed by the MODEL flavour and handed
        to the `codegen` role, which the judge requires verbatim."""
        from wama.common.manifests.codegen import tasks_gen
        from wama.common.tests.tests_dev_ai_bridge import _charger
        codegen = _charger('run_codegen')
        man = {'body': {'params': {'primary': 'PARAMS_JSON', 'schemas': {'PARAMS_JSON': [
            {'name': 'model', 'options_source': 'catalog', 'options_query': {'model_type': 'llm'}}]}},
            'processing': {'model_spec': {'item': {'params_fields': ['prompt', 'model']}}}}}
        preamble = tasks_gen.declared_settings_and_draw(man)
        self.assertIn("opts = effective_settings(_SCH, posees=given)", preamble)
        self.assertIn("resolve_model_choice(opts.pop('model', ''), app_id=ctx.app_id,", preamble)
        self.assertIn("spec={'model_type': 'llm'}", preamble)
        aligned = 'def _process_t(item, ctx):\n' + preamble + '\n    return {"fields": {}}\n'
        self.assertEqual([], codegen.preamble_checks(aligned, '_process_t', preamble))
        # Counter-proof : the glue that read the settings its own way is caught.
        own_way = ('def _process_t(item, ctx):\n'
                   '    opts = effective_settings(PARAMS_JSON, posees={}, preset=item.quality_intent)\n'
                   '    key = resolve_model_choice(item.model, app_id=ctx.app_id, item=item)\n')
        self.assertIn('préambule COMMUN', '\n'.join(
            codegen.preamble_checks(own_way, '_process_t', preamble)))
        # The MODEL flavour composes the very same lines (one source, two readers).
        lines = tasks_gen.settings_and_draw_lines(['a'], param='m', domain={'task': 'ocr'},
                                                  app_id_expr="'reader_01'")
        self.assertEqual('    opts = effective_settings(_SCH, posees=given)', lines[5])
        self.assertIn("app_id='reader_01'", lines[6])

    def test_the_instant_hole_has_its_own_signature(self):
        from wama.common.tests.tests_dev_ai_bridge import _charger
        codegen = _charger('run_codegen')
        glue = 'def instant_result(item):\n    return item.prompt\n'
        self.assertTrue(codegen.controles(glue, 'instant_result',
                                          expected_args=('item',))['signature_ok'])
        # Counter-proof : a task body keeps its (item, ctx).
        self.assertFalse(codegen.controles('def _process_t(item):\n    pass\n',
                                           '_process_t')['signature_ok'])


class _TwinCtx:
    """The task context of `task_skeleton`, reduced to what the twin's two glues use."""

    def __init__(self, item):
        self.item, self.app_id, self.said = item, 'translator_01', []

    def console(self, message):
        self.said.append(message)

    def progress(self, *args):
        pass

    def ensure_model(self, key):
        return key

    def output_place(self, ext, **_kw):
        import os
        from django.conf import settings
        rel = f'users/{self.item.user_id}/translator_01/output/{len(self.said)}_{os.getpid()}{ext}'
        absolute = os.path.join(settings.MEDIA_ROOT, rel)
        os.makedirs(os.path.dirname(absolute), exist_ok=True)
        self.said.append(rel)
        return rel, absolute


class TwinTranslatesAScanTest(TestCase):
    """The REGENERATED twin (`translator_01`, 2026-10-30) end to end, with no engine: a scan
    dropped on the card ticks the OCR switch (`default_from`), a native PDF does not, a posted
    choice wins; then its two glues — the OCR one RENDERED by the generator, the translation
    one written by the `codegen` role — rewrite the scan with the translation as real text."""

    def setUp(self):
        import tempfile
        from django.apps import apps
        from django.test import override_settings
        if not apps.is_installed('wama.translator_01'):
            self.skipTest('jumelle translator_01 absente de ce dépôt')
        self.media = tempfile.mkdtemp()
        override = override_settings(MEDIA_ROOT=self.media)
        override.enable()
        self.addCleanup(override.disable)
        self.addCleanup(shutil.rmtree, self.media, True)
        self.dir = Path(tempfile.mkdtemp(prefix='wama-twin-'))
        self.addCleanup(shutil.rmtree, self.dir, True)
        from wama.common.tests.tests_document_text import LONG, _scan_of
        self.scan, self.ocr = _scan_of(self.dir, [(72, 100, LONG), (72, 114, 'and ends here.')])
        import fitz
        with fitz.open() as pdf:
            pdf.new_page().insert_text((72, 100), LONG)
            pdf.save(str(self.dir / 'native.pdf'))
        from django.contrib.auth import get_user_model
        # Une jumelle ne s'ouvre qu'à son créateur, aux développeurs et aux administrateurs.
        self.user = get_user_model().objects.create_superuser('twin_scan', 'twin@x.fr', 'x')
        self.client.force_login(self.user)

    def _drop(self, path, **post):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from wama.translator_01.models import TranslatorText
        upload = SimpleUploadedFile(path.name, path.read_bytes(), 'application/pdf')
        response = self.client.post('/translator_01/upload/', {'source_document': upload, **post})
        self.assertEqual(200, response.status_code, response.content[:300])
        return TranslatorText.objects.get(pk=response.json()['id'])

    def test_the_switch_follows_the_dropped_file_and_a_posted_choice_wins(self):
        self.assertIs(True, self._drop(self.scan).ocr_scanned_pages)
        self.assertIs(False, self._drop(self.dir / 'native.pdf').ocr_scanned_pages)
        self.assertIs(False, self._drop(self.scan, ocr_scanned_pages='false').ocr_scanned_pages)

    def test_the_two_glues_rewrite_the_scan_with_the_translation(self):
        from types import SimpleNamespace
        import fitz
        from wama.translator_01 import tasks
        item = self._drop(self.scan, target_language='fr')
        engine = SimpleNamespace(read_lines=lambda path: [self.ocr[0]])
        with mock.patch('wama.common.services.ocr_process.positioned_ocr_keys',
                        return_value=['reader:doctr']), \
                mock.patch('wama.common.utils.auto_model.resolve_model_choice',
                           return_value='reader:doctr'), \
                mock.patch('wama.common.backends.manager.backend_instance_for_key',
                           return_value=engine):
            ocr = tasks._process_translate_text_task_ocr(item, _TwinCtx(item))
        item.ocr_file = ocr['fields']['ocr_file']
        item.save(update_fields=['ocr_file'])

        def fake(self, text, *args, **kwargs):
            return {'ok': True, 'text': 'TRADUIT ' + text.split()[0]}
        with mock.patch('wama.common.utils.translator.TranslatorService.translate', fake), \
                mock.patch('wama.common.utils.auto_model.resolve_model_choice',
                           return_value='ollama:translategemma:12b'):
            result = tasks._process_translate_text_task_translate(item, _TwinCtx(item))
        with fitz.open(str(Path(self.media) / result['fields']['output_file'])) as pdf:
            self.assertIn('TRADUIT The', pdf[0].get_text())
