"""The TEXT LAYER of a document (ROADMAP §16.4bis, vision §5.4) — read its passages, rewrite them IN
PLACE keeping everything else ; and the privacy brick's file occultation on top of it.

What is held here, on documents built by the test (in a temporary folder it creates and removes) :
a TXT keeps its encoding ; a DOCX keeps its paragraphs and the formatting of the passage's start,
a name split across two runs, in a table, in a header, or left in a tracked DELETION is gone from
the file ; a native PDF loses the passage's text for good (not just covered) and its metadata ;
the occultation re-reads what it produced and names any original passage still there.
"""
import shutil
import tempfile
import zipfile
from pathlib import Path

from django.test import SimpleTestCase


class _TempDir(SimpleTestCase):

    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix='wama-doctext-'))
        self.addCleanup(shutil.rmtree, self.dir, True)


def _span(segment: str, passage: str) -> tuple:
    start = segment.index(passage)
    return start, start + len(passage)


class PlainTextTest(_TempDir):

    def test_a_cp1252_text_is_rewritten_in_its_own_encoding(self):
        from wama.common.utils.document_text import read_segments, write_replacements
        src = self.dir / 'lettre.txt'
        src.write_bytes('Chère Anne Martin, à bientôt.'.encode('cp1252'))
        [segment] = read_segments(src)
        out = write_replacements(src, self.dir / 'out.txt',
                                 {0: [(*_span(segment, 'Anne Martin'), '[PERSONNE]')]})
        self.assertEqual('Chère [PERSONNE], à bientôt.', Path(out).read_bytes().decode('cp1252'))


class DocxTest(_TempDir):

    def _document(self):
        import docx
        from docx.oxml import parse_xml
        document = docx.Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('Madame ')
        paragraph.add_run('Anne ').bold = True
        paragraph.add_run('Martin').italic = True
        paragraph.add_run(' viendra.')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Responsable : Paul Girard'
        document.sections[0].header.paragraphs[0].text = 'Dossier de Paul Girard'
        # A tracked DELETION : its text stays in the file although no one sees it.
        deleted = parse_xml(
            '<w:del xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'w:id="1" w:author="x"><w:r><w:delText>Jeanne Secret</w:delText></w:r></w:del>')
        document.add_paragraph('Fin.')._p.append(deleted)
        document.core_properties.author = 'Anne Martin'
        path = self.dir / 'note.docx'
        document.save(str(path))
        return path

    def _xml(self, path) -> str:
        with zipfile.ZipFile(path) as z:
            return '\n'.join(z.read(n).decode('utf-8', 'replace') for n in z.namelist()
                             if n.endswith('.xml'))

    def test_names_across_runs_tables_headers_and_deletions_are_gone(self):
        from wama.common.utils.document_text import read_segments, write_replacements
        src = self._document()
        segments = read_segments(src)
        replacements = {}
        for i, segment in enumerate(segments):
            for name, label in (('Anne Martin', '[PERSONNE]'), ('Paul Girard', '[PERSONNE]')):
                if name in segment:
                    replacements.setdefault(i, []).append((*_span(segment, name), label))
        out = Path(write_replacements(src, self.dir / 'out.docx', replacements))
        xml = self._xml(out)
        for gone in ('Anne', 'Martin', 'Paul Girard', 'Jeanne Secret'):
            self.assertNotIn(gone, xml)
        after = read_segments(out)
        self.assertIn('Madame [PERSONNE] viendra.', after)
        self.assertIn('Responsable : [PERSONNE]', after)
        self.assertIn('Dossier de [PERSONNE]', after)
        self.assertEqual(len(segments), len(after))

    def test_the_formatting_of_the_passage_start_is_kept_and_the_rest_untouched(self):
        import docx
        from wama.common.utils.document_text import read_segments, write_replacements
        src = self._document()
        segment = read_segments(src)[0]
        out = write_replacements(src, self.dir / 'out.docx',
                                 {0: [(*_span(segment, 'Anne Martin'), '[PERSONNE]')]})
        runs = docx.Document(out).paragraphs[0].runs
        self.assertEqual(['Madame ', '[PERSONNE]', '', ' viendra.'], [r.text for r in runs])
        self.assertTrue(runs[1].bold)


class PdfTest(_TempDir):

    def _pdf(self):
        import fitz
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((72, 72), 'Contact : Anne Martin, bureau 12.', fontsize=12)
        page.insert_text((72, 100), 'Signature : Anne Martin.', fontsize=12)
        pdf.set_metadata({'author': 'Anne Martin', 'title': 'Dossier Anne Martin'})
        path = self.dir / 'fiche.pdf'
        pdf.save(str(path))
        pdf.close()
        return path

    def test_the_passage_text_is_removed_not_covered_and_the_metadata_scrubbed(self):
        import fitz
        from wama.common.utils.document_text import read_segments, write_replacements
        src = self._pdf()
        [page] = read_segments(src)
        out = write_replacements(src, self.dir / 'out.pdf',
                                 {0: [(*_span(page, 'Anne Martin'), '[PERSONNE]')]})
        with fitz.open(out) as pdf:
            text = pdf[0].get_text()
            self.assertNotIn('Anne Martin', text)          # every occurrence on the page
            self.assertIn('bureau 12', text)               # the rest untouched
            self.assertNotIn('Anne Martin', str(pdf.metadata))
        self.assertNotIn(b'Anne Martin', Path(out).read_bytes())


class PdfBlockFormTest(_TempDir):
    """2026-10-29 (a real translated paper : every block rewritten in sans-serif, no bold, 162
    blocks of 392 shrunk under 80 %) — unit `block` rewrites a block with its OWN dominant style,
    keeps its size by growing into the free room below, keeps its writing direction, and leaves an
    unchanged block untouched."""

    def _spans(self, path):
        import fitz
        with fitz.open(path) as pdf:
            return [s for b in pdf[0].get_text('dict')['blocks'] if b.get('type') == 0
                    for line in b['lines'] for s in line['spans'] if s['text'].strip()]

    def _rewrite(self, src, transform):
        from wama.common.utils.document_text import read_segments, write_replacements
        segments = read_segments(src, unit='block')
        replacements = {i: [(0, len(s), transform(s))] for i, s in enumerate(segments)}
        return write_replacements(src, self.dir / 'out.pdf', replacements, unit='block')

    def test_a_block_keeps_its_family_weight_size_and_color(self):
        import fitz
        src = self.dir / 'in.pdf'
        with fitz.open() as pdf:
            pdf.new_page().insert_text((72, 100), 'Titre en gras', fontname='tibo', fontsize=14,
                                       color=(0, 0, 1))
            pdf.save(str(src))
        [span] = self._spans(self._rewrite(src, str.upper))
        self.assertEqual('TITRE EN GRAS', span['text'].strip())
        self.assertTrue(span['flags'] & 4, 'serif family kept')
        self.assertTrue(span['flags'] & 16, 'bold kept')
        self.assertAlmostEqual(14, span['size'], delta=0.1)
        self.assertEqual(0x0000ff, span['color'])

    def test_a_longer_text_keeps_its_size_by_growing_into_the_free_room_below(self):
        import fitz
        src = self.dir / 'in.pdf'
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((72, 100), 'Un court paragraphe', fontname='tiro', fontsize=11)
            page.insert_text((72, 400), 'Pied', fontname='tiro', fontsize=11)
            pdf.save(str(src))
        longer = lambda s: s if s == 'Pied' else ' '.join([s] * 12)  # noqa: E731
        spans = self._spans(self._rewrite(src, longer))
        body = [s for s in spans if s['text'].strip() != 'Pied']
        self.assertGreater(len({round(s['bbox'][1]) for s in body}), 1, 'wrapped on several lines')
        self.assertTrue(all(abs(s['size'] - 11) < 0.1 for s in body), 'size kept, not shrunk')
        foot = [s for s in spans if s['text'].strip() == 'Pied'][0]
        self.assertLess(max(s['bbox'][3] for s in body), foot['bbox'][1], 'the block below is not covered')

    def test_a_vertical_stamp_is_rewritten_vertical(self):
        import fitz
        src = self.dir / 'in.pdf'
        with fitz.open() as pdf:
            pdf.new_page().insert_text((30, 500), 'Tampon vertical', fontsize=12, rotate=90)
            pdf.save(str(src))
        with fitz.open(self._rewrite(src, str.upper)) as pdf:
            [line] = [line for b in pdf[0].get_text('dict')['blocks'] for line in b['lines']]
        self.assertEqual('TAMPON VERTICAL', ''.join(s['text'] for s in line['spans']).strip())
        self.assertEqual((0, -1), tuple(round(v) for v in line['dir']))

    def test_an_unchanged_block_is_not_rewritten(self):
        """Counter-proof of the font : an untouched block keeps its EMBEDDED font, not the
        generic family a rewrite would give it."""
        import fitz
        src = self.dir / 'in.pdf'
        with fitz.open() as pdf:
            pdf.new_page().insert_text((72, 100), '4.69 1.78', fontname='cour', fontsize=10)
            pdf.save(str(src))
        [span] = self._spans(self._rewrite(src, lambda s: s))
        self.assertEqual('Courier', span['font'])


class OccultFileTest(_TempDir):

    def test_the_occultation_re_reads_its_output_and_names_what_remains(self):
        """Counter-proof : a name detected once but written twice must come back in `remaining`."""
        import docx
        from wama.common.privacy import engine
        from wama.common.utils import detections
        document = docx.Document()
        document.add_paragraph('Écrire à anne@exemple.fr.')
        document.add_paragraph('Anne Martin, encore Anne Martin.')
        src = self.dir / 'lettre.docx'
        document.save(str(src))
        doc = engine.detect_file(src, 'fr', detectors=())          # patterns only : the email
        report = engine.occult_file(src, self.dir / 'out.docx', doc)
        self.assertEqual([], report['remaining'])
        self.assertNotIn('anne@exemple.fr', '\n'.join(p.text for p in docx.Document(report['path']).paragraphs))
        detections.add(doc, 1, [detections.text_span(start=0, end=11, label='PERSON')])
        report = engine.occult_file(src, self.dir / 'out2.docx', doc)
        self.assertEqual(['Anne Martin'], report['remaining'])
        self.assertEqual('[PERSONNE], encore Anne Martin.', docx.Document(report['path']).paragraphs[1].text)
